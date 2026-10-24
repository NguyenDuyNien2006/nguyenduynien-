#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script to generate a professional, academic Microsoft Word (.docx) report
with all actual figures and screenshots embedded directly into the document.
"""

import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

IMAGES_DIR = r"c:\Users\ducshin\.gemini\antigravity-ide\scratch\portfolio-system\docs\images"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for b_name, b_val in borders.items():
        if b_val:
            node = OxmlElement(f'w:{b_name}')
            node.set(qn('w:val'), b_val.get('val', 'single'))
            node.set(qn('w:sz'), str(b_val.get('sz', 4)))
            node.set(qn('w:space'), '0')
            node.set(qn('w:color'), b_val.get('color', 'CCCCCC'))
            tcBorders.append(node)
        else:
            node = OxmlElement(f'w:{b_name}')
            node.set(qn('w:val'), 'none')
            tcBorders.append(node)
    tcPr.append(tcBorders)

def add_embedded_image(doc, img_filename, caption_text, width_cm=15.5):
    """Embeds an actual image file with a centered caption underneath."""
    img_path = os.path.join(IMAGES_DIR, img_filename)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Cm(width_cm))
        
        # Caption
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(10)
        p_cap.paragraph_format.keep_with_next = False
        
        r_bold = p_cap.add_run(caption_text.split(':')[0] + ": ")
        r_bold.bold = True
        r_bold.font.name = "Times New Roman"
        r_bold.font.size = Pt(10.5)
        r_bold.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        
        rest_of_caption = ":".join(caption_text.split(':')[1:]) if ':' in caption_text else ""
        r_text = p_cap.add_run(rest_of_caption)
        r_text.italic = True
        r_text.font.name = "Times New Roman"
        r_text.font.size = Pt(10.5)
        r_text.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    else:
        # Fallback callout box if image not found
        add_callout_box(doc, caption_text)

def add_callout_box(doc, text, title="HÌNH MINH HỌA / CHÚ THÍCH"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Cm(16.0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=160, bottom=160, left=240, right=200)
    set_cell_borders(cell,
                     left={'val': 'single', 'sz': 24, 'color': '2563EB'},
                     top={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                     bottom={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                     right={'val': 'single', 'sz': 4, 'color': 'E2E8F0'})
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    
    r_title = p.add_run(f"📸 {title}: ")
    r_title.bold = True
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(11)
    r_title.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)
    
    r_text = p.add_run(text)
    r_text.italic = True
    r_text.font.name = "Times New Roman"
    r_text.font.size = Pt(11)
    r_text.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_text, lang=""):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Cm(16.0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=160)
    set_cell_borders(cell,
                     left={'val': 'single', 'sz': 18, 'color': '0F172A'},
                     top={'val': 'single', 'sz': 4, 'color': 'CBD5E1'},
                     bottom={'val': 'single', 'sz': 4, 'color': 'CBD5E1'},
                     right={'val': 'single', 'sz': 4, 'color': 'CBD5E1'})
    
    lines = code_text.strip().split("\n")
    for i, line in enumerate(lines):
        if i == 0:
            p = cell.paragraphs[0]
        else:
            p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.05
        
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def format_run_markdown(paragraph, text, base_font_size=13, is_heading=False, heading_color=None):
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(base_font_size)
            if heading_color:
                run.font.color.rgb = heading_color
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(base_font_size)
            if heading_color:
                run.font.color.rgb = heading_color
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(base_font_size - 1)
            run.font.color.rgb = RGBColor(0x99, 0x1B, 0x1B)
        else:
            run = paragraph.add_run(token)
            run.font.name = "Times New Roman"
            run.font.size = Pt(base_font_size)
            if heading_color:
                run.font.color.rgb = heading_color
            if is_heading:
                run.bold = True

def create_document():
    doc = Document()
    
    # 1. Page Margins
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Báo Cáo Thực Hành: Triển khai và Quản trị Hệ thống Phần mềm — Đề tài 13")
        hrun.font.name = "Times New Roman"
        hrun.font.size = Pt(9)
        hrun.font.italic = True
        hrun.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Trang ")
        frun.font.name = "Times New Roman"
        frun.font.size = Pt(10)
        frun.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        f_xml = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        fp._p.append(f_xml)

    # 2. Cover Page
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(0)
    p_inst.paragraph_format.space_after = Pt(2)
    p_inst.paragraph_format.line_spacing = 1.15
    
    r1 = p_inst.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(13)
    r1.bold = True
    
    r2 = p_inst.add_run("ĐẠI HỌC THÁI NGUYÊN\nTRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG\n")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(13.5)
    r2.bold = True
    
    r3 = p_inst.add_run("KHOA CÔNG NGHỆ THÔNG TIN\n")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(13)
    r3.bold = True
    
    r_line = p_inst.add_run("-----------------------***-----------------------")
    r_line.font.name = "Times New Roman"
    r_line.font.size = Pt(11)
    r_line.bold = True
    
    for _ in range(2):
        doc.add_paragraph()
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(6)
    r_main = p_title.add_run("BÁO CÁO THỰC HÀNH TỔNG HỢP\n")
    r_main.font.name = "Times New Roman"
    r_main.font.size = Pt(20)
    r_main.bold = True
    r_main.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(4)
    p_sub.paragraph_format.space_after = Pt(16)
    r_sub = p_sub.add_run("MÔN HỌC: TRIỂN KHAI VÀ QUẢN TRỊ HỆ THỐNG PHẦN MỀM")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(14.5)
    r_sub.bold = True
    
    tbl_topic = doc.add_table(rows=1, cols=1)
    tbl_topic.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_topic = tbl_topic.cell(0, 0)
    c_topic.width = Cm(15.5)
    set_cell_background(c_topic, "F1F5F9")
    set_cell_margins(c_topic, top=200, bottom=200, left=240, right=240)
    set_cell_borders(c_topic,
                     top={'val': 'single', 'sz': 12, 'color': '1B365D'},
                     bottom={'val': 'single', 'sz': 12, 'color': '1B365D'},
                     left={'val': 'single', 'sz': 12, 'color': '1B365D'},
                     right={'val': 'single', 'sz': 12, 'color': '1B365D'})
    
    p_top_inner = c_topic.paragraphs[0]
    p_top_inner.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_top_inner.paragraph_format.line_spacing = 1.25
    r_de = p_top_inner.add_run("ĐỀ TÀI 13:\n")
    r_de.font.name = "Times New Roman"
    r_de.font.size = Pt(14)
    r_de.bold = True
    r_de.font.color.rgb = RGBColor(0x0F, 0x4C, 0x81)
    
    r_detai = p_top_inner.add_run(
        "TRIỂN KHAI VÀ QUẢN TRỊ HỆ THỐNG WEBSITE PORTFOLIO CÁ NHÂN VỚI "
        "NGINX REVERSE PROXY, DOCKER COMPOSE, MYSQL, PHPMYADMIN, "
        "PROMETHEUS, GRAFANA, LOKI VÀ HARDENING BẢO MẬT"
    )
    r_detai.font.name = "Times New Roman"
    r_detai.font.size = Pt(14)
    r_detai.bold = True
    r_detai.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    for _ in range(2):
        doc.add_paragraph()
        
    info_table = doc.add_table(rows=5, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("Giảng viên hướng dẫn:", "ThS. / TS. Giảng viên phụ trách môn học"),
        ("Sinh viên thực hiện:", "Nguyễn Duy Niên"),
        ("Mã số sinh viên (MSSV):", "DTC245200301"),
        ("Lớp / Khóa:", "K23 - TKQT HTPM"),
        ("Học kỳ / Năm học:", "Học kỳ 1 - Năm học 2026")
    ]
    
    for row_idx, (label, val) in enumerate(info_data):
        row = info_table.rows[row_idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Cm(6.5)
        c1.width = Cm(8.5)
        set_cell_margins(c0, top=100, bottom=100, left=100, right=100)
        set_cell_margins(c1, top=100, bottom=100, left=100, right=100)
        set_cell_borders(c0)
        set_cell_borders(c1)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.line_spacing = 1.2
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(label)
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(12.5)
        r0.bold = True
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.line_spacing = 1.2
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(val)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(12.5)
        if label in ["Sinh viên thực hiện:", "Mã số sinh viên (MSSV):"]:
            r1.bold = True
            
    for _ in range(3):
        doc.add_paragraph()
        
    p_bot = doc.add_paragraph()
    p_bot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bot.paragraph_format.space_before = Pt(10)
    p_bot.paragraph_format.space_after = Pt(0)
    r_bot = p_bot.add_run("THÁI NGUYÊN, THÁNG 10 NĂM 2026")
    r_bot.font.name = "Times New Roman"
    r_bot.font.size = Pt(12)
    r_bot.bold = True
    
    doc.add_page_break()
    
    # 3. Read and Parse BAO_CAO_TKQT_HTPM.md
    md_path = r"c:\Users\ducshin\.gemini\antigravity-ide\scratch\portfolio-system\docs\BAO_CAO_TKQT_HTPM.md"
    with open(md_path, "r", encoding="utf-8") as f:
        md_content = f.read()
        
    lines = md_content.splitlines()
    
    in_code_block = False
    code_block_lang = ""
    code_block_lines = []
    
    in_table = False
    table_lines = []
    
    skip_cover_html = True
    
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        
        if skip_cover_html:
            if stripped == "# MỤC LỤC CHI TIẾT":
                skip_cover_html = False
            else:
                i += 1
                continue
                
        # Code block handling
        if stripped.startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_block_lang = stripped[3:].strip()
                code_block_lines = []
            else:
                in_code_block = False
                # If this was mermaid in Chapter 2, replace with the beautiful Architecture diagram!
                if code_block_lang == "mermaid":
                    add_embedded_image(
                        doc,
                        "hinh_2_1_architecture.png",
                        "Hình 2.1: Sơ đồ kiến trúc tổng thể hệ thống đa tầng và phân vùng mạng Docker (3 Networks)",
                        width_cm=16.0
                    )
                else:
                    add_code_block(doc, "\n".join(code_block_lines), code_block_lang)
                code_block_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue
            
        # Table handling
        if stripped.startswith("|") and stripped.endswith("|"):
            if not in_table:
                in_table = True
                table_lines = [stripped]
            else:
                table_lines.append(stripped)
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                parse_and_add_table(doc, table_lines)
                table_lines = []
                
        if not stripped:
            i += 1
            continue
            
        if "page-break-after: always;" in stripped:
            doc.add_page_break()
            i += 1
            continue
            
        if stripped == "---":
            i += 1
            continue
            
        # Image placeholder replacements with ACTUAL EMBEDDED IMAGES!
        if stripped.startswith("*(") and stripped.endswith(")*"):
            caption_text = stripped[2:-2]
            
            if "3.1" in caption_text:
                add_embedded_image(doc, "hinh_3_1_phpmyadmin.png", "Hình 3.1: Giao diện cấu trúc các bảng CSDL portfolio_db trên phpMyAdmin")
            elif "3.2" in caption_text:
                add_embedded_image(doc, "hinh_3_2_portfolio_cms.png", "Hình 3.2: Giao diện Website Portfolio Dark Glassmorphism và Bảng điều khiển Quản trị Admin CMS")
            elif "3.3" in caption_text:
                add_embedded_image(doc, "hinh_3_3_phpmyadmin_browse.png", "Hình 3.3: Giao diện đăng nhập và duyệt bảng dữ liệu trong phpMyAdmin")
            elif "4.1" in caption_text:
                add_embedded_image(doc, "hinh_4_1_nginx_ssl_headers.png", "Hình 4.1: Kiểm tra chứng chỉ SSL tự ký và trọn bộ Security Headers trên Nginx")
            elif "5.1" in caption_text:
                add_embedded_image(doc, "hinh_5_1_grafana_dashboard.png", "Hình 5.1: Dashboard giám sát tổng thể hiệu năng và số liệu ứng dụng trên Grafana")
            elif "6.1" in caption_text:
                add_embedded_image(doc, "hinh_6_1_loki_logql.png", "Hình 6.1: Giao diện Grafana Explore thực thi truy vấn LogQL thu thập nhật ký với Loki")
            elif "7.1" in caption_text:
                add_embedded_image(doc, "hinh_7_1_hardening_verify.png", "Hình 7.1: Kiểm chứng các biện pháp Hardening bảo mật (Non-root user, Network isolation, Resource limits)")
            elif "8.1" in caption_text or "8.6" in caption_text:
                # Add test scenario screenshots
                add_embedded_image(doc, "hinh_8_1_docker_compose_ps.png", "Hình 8.1: Kết quả kiểm thử Kịch bản 1 — Khởi động 11 containers thành công với Docker Compose")
                add_embedded_image(doc, "hinh_8_4_grafana_verify.png", "Hình 8.4: Kết quả kiểm thử Kịch bản 4 — Giám sát luồng truy cập và tài nguyên trên Grafana")
                add_embedded_image(doc, "hinh_8_5_loki_verify.png", "Hình 8.5: Kết quả kiểm thử Kịch bản 5 — Phân tích lỗi HTTP >= 400 bằng truy vấn LogQL")
                add_embedded_image(doc, "hinh_7_1_hardening_verify.png", "Hình 8.6: Kết quả kiểm thử Kịch bản 6 — Kiểm chứng an toàn: tài khoản 'node' và cổng 3306 ẩn giấu")
            else:
                add_callout_box(doc, caption_text)
                
            i += 1
            continue
            
        # Headings
        if stripped.startswith("# "):
            heading_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            format_run_markdown(p, heading_text, base_font_size=16, is_heading=True, heading_color=RGBColor(0x1B, 0x36, 0x5D))
            i += 1
            continue
            
        if stripped.startswith("## "):
            heading_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(13)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.keep_with_next = True
            format_run_markdown(p, heading_text, base_font_size=14, is_heading=True, heading_color=RGBColor(0x0F, 0x4C, 0x81))
            i += 1
            continue
            
        if stripped.startswith("### "):
            heading_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(9)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            format_run_markdown(p, heading_text, base_font_size=13, is_heading=True, heading_color=RGBColor(0x1F, 0x29, 0x37))
            i += 1
            continue
            
        # In Chapter 9, add git commit image after the commit list description
        if "tệp tài liệu hướng dẫn README chi tiết" in stripped:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.25
            format_run_markdown(p, stripped, base_font_size=13)
            
            # Embed Git commit screenshot
            add_embedded_image(doc, "hinh_9_1_git_commits.png", "Hình 9.1: Lịch sử Git Log thể hiện 03 commits có ý nghĩa tương ứng với 3 giai đoạn hoàn thiện đề tài")
            i += 1
            continue
            
        # Bullet list items
        if stripped.startswith("- ") or stripped.startswith("* "):
            bullet_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
            format_run_markdown(p, bullet_text, base_font_size=13)
            i += 1
            continue
            
        # Numbered list items
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            num = num_match.group(1)
            num_text = num_match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
            format_run_markdown(p, num_text, base_font_size=13)
            i += 1
            continue
            
        # Sign-off block
        if "<div align=\"center\">" in stripped or "NGƯỜI LẬP BÁO CÁO" in stripped:
            add_signoff_section(doc)
            while i < len(lines) and "</div>" not in lines[i]:
                i += 1
            i += 1
            continue
            
        # Regular paragraph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.25
        format_run_markdown(p, stripped, base_font_size=13)
        
        i += 1
        
    if in_table and table_lines:
        parse_and_add_table(doc, table_lines)
        
    out_dir = r"c:\Users\ducshin\.gemini\antigravity-ide\scratch\portfolio-system\docs"
    
    # Primary target file: BAO_CAO_TKQT_HTPM_HOAN_CHINH.docx
    target_names = [
        "BAO_CAO_TKQT_HTPM_HOAN_CHINH.docx",
        "BAO_CAO_TKQT_HTPM_CO_ANH.docx",
        "BAO_CAO_TKQT_HTPM.docx"
    ]
    
    saved_files = []
    for t_name in target_names:
        p = os.path.join(out_dir, t_name)
        try:
            doc.save(p)
            saved_files.append(p)
            print(f"Successfully saved: {p}")
        except PermissionError:
            print(f"File {t_name} is currently open in Word. Skipped.")
            
    print("Document export completed.")

def parse_and_add_table(doc, table_lines):
    if len(table_lines) < 2:
        return
    rows_data = []
    for line in table_lines:
        if re.match(r'^\|[\s:\-]+\|', line) and not any(c.isalnum() for c in line):
            continue
        cells = [c.strip() for c in line.strip('|').split('|')]
        rows_data.append(cells)
    if not rows_data:
        return
        
    num_cols = max(len(r) for r in rows_data)
    tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    total_width_cm = 16.0
    col_w_cm = total_width_cm / num_cols
    
    for row_idx, row in enumerate(rows_data):
        is_header = (row_idx == 0)
        table_row = tbl.rows[row_idx]
        for col_idx in range(num_cols):
            cell = table_row.cells[col_idx]
            cell.width = Cm(col_w_cm)
            cell_text = row[col_idx] if col_idx < len(row) else ""
            if is_header:
                set_cell_background(cell, "1B365D")
                set_cell_margins(cell, top=140, bottom=140, left=120, right=120)
                set_cell_borders(cell,
                                 top={'val': 'single', 'sz': 8, 'color': '1B365D'},
                                 bottom={'val': 'single', 'sz': 12, 'color': '0F4C81'},
                                 left={'val': 'single', 'sz': 4, 'color': '3B82F6'},
                                 right={'val': 'single', 'sz': 4, 'color': '3B82F6'})
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(cell_text.replace('**', ''))
                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            else:
                bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                set_cell_borders(cell,
                                 top={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                                 bottom={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                                 left={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                                 right={'val': 'single', 'sz': 4, 'color': 'E2E8F0'})
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                if col_idx in [0, 3, 4] and len(cell_text) < 15:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                format_run_markdown(p, cell_text, base_font_size=11)
                
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_signoff_section(doc):
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(12)
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c0 = tbl.cell(0, 0)
    c1 = tbl.cell(0, 1)
    c0.width = Cm(8.0)
    c1.width = Cm(8.0)
    set_cell_borders(c0)
    set_cell_borders(c1)
    
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.line_spacing = 1.2
    
    r_loc = p1.add_run("Thái Nguyên, ngày 23 tháng 10 năm 2026\n")
    r_loc.font.name = "Times New Roman"
    r_loc.font.size = Pt(12)
    r_loc.font.italic = True
    
    r_role = p1.add_run("NGƯỜI LẬP BÁO CÁO\n")
    r_role.font.name = "Times New Roman"
    r_role.font.size = Pt(13)
    r_role.bold = True
    
    r_note = p1.add_run("(Ký và ghi rõ họ tên)\n\n\n\n\n")
    r_note.font.name = "Times New Roman"
    r_note.font.size = Pt(11)
    r_note.font.italic = True
    
    r_name = p1.add_run("Nguyễn Duy Niên")
    r_name.font.name = "Times New Roman"
    r_name.font.size = Pt(13.5)
    r_name.bold = True

if __name__ == "__main__":
    create_document()
